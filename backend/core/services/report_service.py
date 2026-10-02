"""
Report and Export Service for GoldenSwan Hotel.
Generates comprehensive report statistics, formatted Excel workbooks (via openpyxl),
and styled PDF documents (via reportlab).
"""

import io
from typing import Dict, Any, List, Optional
from datetime import datetime, timezone
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from core.database import get_database
from core.repositories import (
    payments_repo,
    bookings_repo,
    expenses_repo,
    site_expenses_repo,
    staff_repo,
    attendance_repo,
    payroll_repo,
)
from core.serializers.mongodb import mongo_to_json


class ReportService:
    def get_financial_summary(self, start_date: Optional[datetime] = None, end_date: Optional[datetime] = None) -> Dict[str, Any]:
        db = get_database()
        query = {}
        if start_date and end_date:
            query = {"payment_date": {"$gte": start_date, "$lt": end_date}}

        # Revenue
        rev_pipe = []
        if query:
            rev_pipe.append({"$match": query})
        rev_pipe.append({"$group": {"_id": None, "total": {"$sum": "$amount"}}})
        rev_res = list(db["payments"].aggregate(rev_pipe))
        total_revenue = float(str(rev_res[0]["total"])) if rev_res and "total" in rev_res[0] else 0.0

        # Expenses
        exp_match = {}
        if start_date and end_date:
            exp_match = {"date": {"$gte": start_date, "$lt": end_date}}
        exp_pipe = []
        if exp_match:
            exp_pipe.append({"$match": exp_match})
        exp_pipe.append({"$group": {"_id": None, "total": {"$sum": "$total_amount"}}})
        exp_res = list(db["expenses"].aggregate(exp_pipe))
        total_expenses = float(str(exp_res[0]["total"])) if exp_res and "total" in exp_res[0] else 0.0

        # Site expenses
        site_match = {}
        if start_date and end_date:
            site_match = {"date": {"$gte": start_date, "$lt": end_date}}
        site_pipe = []
        if site_match:
            site_pipe.append({"$match": site_match})
        site_pipe.append({"$group": {"_id": None, "total": {"$sum": "$total_amount"}}})
        site_res = list(db["site_expenses"].aggregate(site_pipe))
        total_site_expenses = float(str(site_res[0]["total"])) if site_res and "total" in site_res[0] else 0.0

        net_operating_amount = total_revenue - (total_expenses + total_site_expenses)

        return {
            "total_revenue": round(total_revenue, 2),
            "total_expenses": round(total_expenses, 2),
            "total_site_expenses": round(total_site_expenses, 2),
            "net_operating_amount": round(net_operating_amount, 2),
        }

    def export_excel(self, module: str, start_date: Optional[datetime] = None, end_date: Optional[datetime] = None) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = module.capitalize()

        # Styles
        title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep navy
        border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # Header Title
        ws.merge_cells("A1:G1")
        ws["A1"] = f"GoldenSwan Hotel - {module.upper()} REPORT"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:G2")
        ws["A2"] = f"Generated on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
        ws["A2"].alignment = Alignment(horizontal="center")

        row_start = 4

        if module == "bookings":
            headers = ["Booking ID", "Customer", "Phone", "Room", "Check-In", "Check-Out", "Total (INR)", "Paid (INR)", "Balance (INR)", "Status"]
            bookings = bookings_repo.find({}, sort=[("created_at", -1)], limit=500)
            ws.append([])
            ws.append(headers)

            for b in bookings:
                ws.append([
                    b.get("booking_id", ""),
                    b.get("customer_name", ""),
                    b.get("customer_phone", ""),
                    b.get("room_number", ""),
                    b.get("check_in").strftime("%Y-%m-%d") if isinstance(b.get("check_in"), datetime) else str(b.get("check_in", "")),
                    b.get("check_out").strftime("%Y-%m-%d") if isinstance(b.get("check_out"), datetime) else str(b.get("check_out", "")),
                    float(str(b.get("total_amount", "0"))),
                    float(str(b.get("amount_paid", "0"))),
                    float(str(b.get("balance_amount", "0"))),
                    b.get("booking_status", ""),
                ])

        elif module == "expenses":
            headers = ["Expense ID", "Date", "Category", "Description", "Total (INR)", "Paid (INR)", "Balance (INR)", "Payment Mode", "Status"]
            expenses = expenses_repo.find({}, sort=[("date", -1)], limit=500)
            ws.append([])
            ws.append(headers)

            for e in expenses:
                ws.append([
                    e.get("expense_id", ""),
                    e.get("date").strftime("%Y-%m-%d") if isinstance(e.get("date"), datetime) else str(e.get("date", "")),
                    e.get("category_name", ""),
                    e.get("description", ""),
                    float(str(e.get("total_amount", "0"))),
                    float(str(e.get("amount_paid", "0"))),
                    float(str(e.get("balance", "0"))),
                    e.get("payment_mode", ""),
                    e.get("payment_status", ""),
                ])

        else: # payments
            headers = ["Payment ID", "Booking ID", "Customer", "Date", "Amount (INR)", "Mode", "Reference", "Received By"]
            payments = payments_repo.find({}, sort=[("payment_date", -1)], limit=500)
            ws.append([])
            ws.append(headers)

            for p in payments:
                ws.append([
                    p.get("payment_id", ""),
                    p.get("booking_code", ""),
                    p.get("customer_name", ""),
                    p.get("payment_date").strftime("%Y-%m-%d") if isinstance(p.get("payment_date"), datetime) else str(p.get("payment_date", "")),
                    float(str(p.get("amount", "0"))),
                    p.get("payment_mode", ""),
                    p.get("reference_number", ""),
                    p.get("received_by", ""),
                ])

        # Style table headers
        for cell in ws[row_start]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def export_pdf(self, module: str) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            textColor=colors.HexColor("#1E3A8A"),
            spaceAfter=6,
        )
        subtitle_style = ParagraphStyle(
            "SubTitleStyle",
            parent=styles["Normal"],
            fontSize=9,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=15,
        )

        elements.append(Paragraph("<b>GoldenSwan Hotel</b>", title_style))
        elements.append(Paragraph(f"Official {module.capitalize()} Summary &bull; Generated on {datetime.now(timezone.utc).strftime('%d %b %Y, %H:%M UTC')}", subtitle_style))
        elements.append(Spacer(1, 10))

        if module == "bookings":
            bookings = bookings_repo.find({}, sort=[("created_at", -1)], limit=30)
            data = [["Booking ID", "Customer", "Room", "Check-in", "Check-out", "Total", "Status"]]
            for b in bookings:
                data.append([
                    str(b.get("booking_id", "")),
                    str(b.get("customer_name", ""))[:18],
                    str(b.get("room_number", "")),
                    b.get("check_in").strftime("%d-%m-%Y") if isinstance(b.get("check_in"), datetime) else "",
                    b.get("check_out").strftime("%d-%m-%Y") if isinstance(b.get("check_out"), datetime) else "",
                    f"INR {float(str(b.get('total_amount', '0'))):,.2f}",
                    str(b.get("booking_status", "")),
                ])
        else: # expenses
            expenses = expenses_repo.find({}, sort=[("date", -1)], limit=30)
            data = [["Expense ID", "Category", "Description", "Date", "Total", "Paid", "Status"]]
            for e in expenses:
                data.append([
                    str(e.get("expense_id", "")),
                    str(e.get("category_name", ""))[:14],
                    str(e.get("description", ""))[:20],
                    e.get("date").strftime("%d-%m-%Y") if isinstance(e.get("date"), datetime) else "",
                    f"INR {float(str(e.get('total_amount', '0'))):,.2f}",
                    f"INR {float(str(e.get('amount_paid', '0'))):,.2f}",
                    str(e.get("payment_status", "")),
                ])

        t = Table(data)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        elements.append(t)

        elements.append(Spacer(1, 20))
        elements.append(Paragraph("<font size='8' color='#94A3B8'>GoldenSwan Hotel Management System &bull; Confidential &bull; Page 1 of 1</font>", styles["Normal"]))

        doc.build(elements)
        buffer.seek(0)
        return buffer


report_service = ReportService()
